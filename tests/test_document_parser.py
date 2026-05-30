from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from backend.core.document_parser import extract_text


def test_extract_text_from_plain_text_file() -> None:
    """Plain text uploads are normalized and returned without external services."""
    parsed = extract_text(BytesIO(b"Python   SQL\ncareer plan"), "resume.txt")
    assert parsed.text == "Python SQL\ncareer plan"
    assert parsed.extension == ".txt"


def test_extract_text_from_docx_file() -> None:
    """DOCX uploads are parsed through python-docx from a file-like object."""
    document = Document()
    document.add_paragraph("Built machine learning dashboards.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    parsed = extract_text(buffer, "resume.docx")

    assert "Built machine learning dashboards." in parsed.text
    assert "Python" in parsed.text
    assert parsed.metadata["paragraphs"] == 1


def test_extract_text_from_pdf_file() -> None:
    """PDF uploads are parsed through pdfplumber from a file-like object."""
    buffer = BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, "Resume with Python and SQL projects")
    page.save()
    buffer.seek(0)

    parsed = extract_text(buffer, "resume.pdf")

    assert "Python" in parsed.text
    assert parsed.metadata["pages"] == 1


def test_unsupported_file_extension_raises() -> None:
    """Unsupported uploads fail with a clear validation error."""
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(BytesIO(b"content"), "resume.exe")
